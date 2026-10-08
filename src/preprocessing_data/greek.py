"""Add the Esteban and labelled Greek datasets to TrOCR pretraining.

The source datasets are added to both ``greek/pretraining`` and
``combined/pretraining``. Finetuning partitions are deliberately untouched.
Running the script repeatedly is safe: previously imported rows are replaced.
"""

from __future__ import annotations

import json
import os
import random
import re
import shutil
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
PROCESSED_ROOT = REPO_ROOT / "data" / "processed"
ESTEBAN_ROOT = REPO_ROOT / "data" / "raw" / "greek_estaban"
LABELLED_ROOT = REPO_ROOT / "data" / "raw" / "greek_labelled_data"
TARGETS = (
    PROCESSED_ROOT / "greek" / "pretraining",
    PROCESSED_ROOT / "combined" / "pretraining",
)
SPLITS = ("train", "val", "test")
IMPORTED_PREFIXES = ("esteban__", "labelled__")
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
SPLIT_SEED = 1111

# Keyboard lookalikes. Latin ß is mapped separately onto stigma.
LATIN_LOOKALIKE_TO_GREEK = str.maketrans(
    {
        "A": "Α",
        "E": "Ε",
        "H": "Η",
        "I": "Ι",
        "K": "Κ",
        "M": "Μ",
        "N": "Ν",
        "O": "Ο",
        "T": "Τ",
        "X": "Χ",
        "Y": "Υ",
        "o": "ο",
        "t": "τ",
    }
)


def normalize_latin_lookalikes(text: str) -> str:
    """Map leftover Latin keys onto the Greek letters they look like."""
    return text.translate(LATIN_LOOKALIKE_TO_GREEK)


LATIN_SHARP_S = "\u00df"  # ß
GREEK_STIGMA = "\u03db"  # ϛ
SHARP_S_TO_STIGMA = str.maketrans({LATIN_SHARP_S: GREEK_STIGMA})


def normalize_stigma(text: str) -> str:
    """Map Latin sharp S onto Greek stigma.

    Eparchos section numbers use ß as the printed stand-in for numeral 6.
    """
    return text.translate(SHARP_S_TO_STIGMA)


MIDDLE_DOT = "\u00b7"  # ·
GREEK_ANO_TELEIA = "\u0387"  # ·
BULLET = "\u2022"  # •
RAISED_DOTS_TO_ANO_TELEIA = str.maketrans({MIDDLE_DOT: GREEK_ANO_TELEIA, BULLET: GREEK_ANO_TELEIA})


def normalize_raised_dots(text: str) -> str:
    """Map middle-dot and bullet onto Greek ano teleia.

    ASCII period is left as-is: it may be a low full stop. Do not NFC after
    this step — Unicode decomposes U+0387 to U+00B7.
    """
    return text.translate(RAISED_DOTS_TO_ANO_TELEIA)


GREEK_KORONIS = "\u1fbd"  # ᾽
GREEK_PSILI = "\u1fbf"  # ᾿
RIGHT_SINGLE_QUOTE = "\u2019"  # ’
ASCII_APOSTROPHE = "'"
MODIFIER_APOSTROPHE = "\u02bc"  # ʼ
HIGH_COMMAS_TO_KORONIS = str.maketrans(
    {
        GREEK_PSILI: GREEK_KORONIS,
        RIGHT_SINGLE_QUOTE: GREEK_KORONIS,
        ASCII_APOSTROPHE: GREEK_KORONIS,
        MODIFIER_APOSTROPHE: GREEK_KORONIS,
    }
)


def normalize_high_commas(text: str) -> str:
    """Map lookalike elision/breathing commas onto Greek koronis.

    Spacing psili, Unicode/ASCII apostrophes, and modifier apostrophe are
    the same raised comma in this GT as koronis.
    """
    return text.translate(HIGH_COMMAS_TO_KORONIS)


ACUTE_ACCENT = "\u00b4"  # ´
GREEK_OXIA = "\u1ffd"  # ´
GREEK_TONOS = "\u0384"  # ΄
GREEK_NUMERAL_SIGN = "\u0374"  # ʹ
TICKS_TO_KORONIS = str.maketrans(
    {
        ACUTE_ACCENT: GREEK_KORONIS,
        GREEK_OXIA: GREEK_KORONIS,
        GREEK_TONOS: GREEK_KORONIS,
        GREEK_NUMERAL_SIGN: GREEK_KORONIS,
    }
)


def normalize_apostrophe_ticks(text: str) -> str:
    """Map acute / oxia / tonos / keraia onto Greek koronis.

    On the page these are one tick. Meaning (number vs elision) comes from
    context. None of the source marks is an apostrophe; koronis is.
    """
    return text.translate(TICKS_TO_KORONIS)


FOUR_DOT_PUNCTUATION = "\u2058"  # ⁘
FOUR_DOT_MARK = "\u205b"  # ⁛
FOUR_DOTS_TO_PUNCTUATION = str.maketrans({FOUR_DOT_MARK: FOUR_DOT_PUNCTUATION})


def normalize_four_dots(text: str) -> str:
    """Map four-dot mark onto four-dot punctuation."""
    return text.translate(FOUR_DOTS_TO_PUNCTUATION)


ASCII_SEMICOLON = ";"
GREEK_QUESTION_MARK = "\u037e"  # ;
SEMICOLON_TO_EROTIMATIKO = str.maketrans({ASCII_SEMICOLON: GREEK_QUESTION_MARK})


def normalize_question_mark(text: str) -> str:
    """Map ASCII semicolon onto Greek question mark.

    Same ink. Do not NFC after this step — Unicode decomposes U+037E to U+003B.
    """
    return text.translate(SEMICOLON_TO_EROTIMATIKO)


HALF_TRIANGULAR_COLON = "\u02d1"  # ˑ
ASCII_COLON = ":"
COLON_LOOKALIKES_TO_COLON = str.maketrans({HALF_TRIANGULAR_COLON: ASCII_COLON})


def normalize_colon(text: str) -> str:
    """Map the manuscript half-colon onto the two-dot colon.

    `:ˑ` becomes `::` after the map; collapse that to one colon.
    """
    text = text.translate(COLON_LOOKALIKES_TO_COLON)
    while "::" in text:
        text = text.replace("::", ASCII_COLON)
    return text


EN_DASH = "\u2013"  # –
EM_DASH = "\u2014"  # —
ASCII_HYPHEN = "-"
DASHES_TO_HYPHEN = str.maketrans({EN_DASH: ASCII_HYPHEN, EM_DASH: ASCII_HYPHEN})


def normalize_dash(text: str) -> str:
    """Map en/em dashes onto ASCII hyphen."""
    return text.translate(DASHES_TO_HYPHEN)


EDITORIAL_LETTER_COUNT = re.compile(r"\[[0-9]+\]")
MULTIPLE_SPACES = re.compile(r" {2,}")
GREEK_KAI_SYMBOL = "\u03d7"  # ϗ
COMBINING_OVERLINE = "\u0305"  # ̅, numeral bar
SPACING_MACRON = "\u00af"  # ¯
GREEK_STIGMA_NUMERAL = "\u03db"  # ϛ, numeral 6
DIGIT_TO_GREEK_NUMERAL = {
    "1": "α",
    "2": "β",
    "3": "γ",
    "4": "δ",
    "5": "ε",
    "6": GREEK_STIGMA_NUMERAL,
    "7": "ζ",
    "8": "η",
    "9": "θ",
}
KAI_STROKE_TOKEN = re.compile(r"<\+\+>")
KAI_TOKEN = re.compile(r"<\+>")
NUMERAL_TOKEN = re.compile(r"<([1-9])>")
DELETED_TOKEN = re.compile(r"<->")
SPACED_MACRON = re.compile(r"\s*" + SPACING_MACRON)
DRESDEN_PREFIX = "dresden_a151_facs"
CPGR23_PREFIX = "greek_htr_cpgr23"
# Modern convention: σ inside a word, ς at its end. A σ (with its marks) is
# word-final before whitespace, punctuation/symbols or the line end.
WORD_FINAL_SIGMA = re.compile(r"σ(?P<marks>[\u0300-\u036f]*)(?=$|[^\w\u0300-\u036f])")
# ς (with its marks) directly followed by a letter: a missing word space.
FINAL_SIGMA_BEFORE_LETTER = re.compile(r"ς[\u0300-\u036f]*(?=[^\W\d_])")
# ἱςρόν is a typo for ἱερόν, not a missing space: left as transcribed.
SIGMA_SPACE_EXCLUDED_IMAGES = frozenset(
    {"greek_hpgtr_0007636.jpg", "esteban__bodleian_14_page_037_line_018.png"}
)
EPARCHOS_PREFIX = "greek_eparchos"
# Line-end hyphen after a letter, or after spaces (removed with them). Never
# after ":" -- the ":-" colon-dash is a drawn flourish and stays.
EPARCHOS_LINE_END_HYPHEN = re.compile(r"(?<=[^\W\d_])-$|\s+-$")
# One row has "ἐδωδί- -ἀναρρώνυσιν": both stray hyphens go, one space stays.
EPARCHOS_DOUBLE_HYPHEN = re.compile(r"(?<=[^\W\d_])- -(?=[^\W\d_])")
# A letter plus any combining marks already on it.
LETTER_WITH_MARKS = r"[^\W\d_][\u0300-\u036f\u1dc0-\u1dff\u20d0-\u20ff\ufe20-\ufe2f]*"
# Dresden editorial expansion "(...)" with the letters flanking it, if any.
FLANKED_EXPANSION = re.compile(
    rf"(?P<before>{LETTER_WITH_MARKS})?\([^()\s]*\)(?P<after>{LETTER_WITH_MARKS})?"
)


def bar_letter(letter: str | None) -> str:
    """Append U+0305 after a letter and its marks, unless it is already barred."""
    if not letter:
        return ""
    return letter if COMBINING_OVERLINE in letter else letter + COMBINING_OVERLINE


def remove_expansion_bar_flanks(match: re.Match[str]) -> str:
    """Delete one expansion and bar the letters drawn on each side of it.

    The scribe's bar covers the contracted letters: θ(εὸ)ς -> θ̅ς̅,
    στ(αυ)ροῦ -> στ̅ρ̅οῦ, προκ(είμενον) -> προκ̅.
    """
    return bar_letter(match.group("before")) + bar_letter(match.group("after"))


# Editorial GT tokens resolved to page ink, applied in this order by
# ``apply_editorial_rules``. Each entry is (report name, pattern, replacement,
# rows: an image-name prefix, a predicate on the image name, or None for every
# Greek row).
EDITORIAL_RULES: tuple[
    tuple[str, re.Pattern[str], object, str | Callable[[str], bool] | None], ...
] = (
    # Stavronikita marks an illegible stretch with the number of missing
    # letters in square brackets. That number is not ink: remove it.
    ("[N] editorial letter count -> removed", EDITORIAL_LETTER_COUNT, "", None),
    # <++>: the kappa-with-stroke form of the kai abbreviation; an allograph
    # of the same character as <+>.
    ("<++> -> U+03D7 kai symbol", KAI_STROKE_TOKEN, GREEK_KAI_SYMBOL, None),
    # <+>: the S-shaped Byzantine kai abbreviation drawn in the crop.
    ("<+> -> U+03D7 kai symbol", KAI_TOKEN, GREEK_KAI_SYMBOL, None),
    # <n>: the ink is the Greek numeral letter with a bar above it.
    (
        "<n> -> Greek numeral letter + U+0305",
        NUMERAL_TOKEN,
        lambda match: DIGIT_TO_GREEK_NUMERAL[match.group(1)] + COMBINING_OVERLINE,
        None,
    ),
    # <->: nothing distinct in the ink; delete.
    ("<-> -> removed", DELETED_TOKEN, "", None),
    # Numeral bars have one encoding: spacing macron becomes a combining
    # overline on the preceding letter (dresden "ιγ¯" -> "ιγ̅").
    ("U+00AF -> U+0305 on preceding letter", SPACED_MACRON, COMBINING_OVERLINE, None),
    # Eparchos: the ink shows no hyphen at the line end.
    ("Eparchos letter- - letter -> letter letter", EPARCHOS_DOUBLE_HYPHEN, " ", EPARCHOS_PREFIX),
    ("Eparchos line-end hyphen -> removed", EPARCHOS_LINE_END_HYPHEN, "", EPARCHOS_PREFIX),
    # Dresden nomina sacra: the editor's expansion in (...) was never drawn;
    # keep only the ink, with the scribe's bar on the flanking letters.
    # Dresden only: the printed New Testament has real parentheses.
    (
        "Dresden (expansion) -> removed, flanking letters + U+0305",
        FLANKED_EXPANSION,
        remove_expansion_bar_flanks,
        DRESDEN_PREFIX,
    ),
    # cpgr23 writes word-final sigma as σ; other sources' final σ are line
    # splits (κτίσ¬), elisions (φήσ᾽) or abbreviations and stay.
    ("cpgr23 word-final σ -> ς", WORD_FINAL_SIGMA, r"ς\g<marks>", CPGR23_PREFIX),
    (
        "ς + letter -> ς + space (missing word space)",
        FINAL_SIGMA_BEFORE_LETTER,
        lambda match: match.group(0) + " ",
        lambda image_name: image_name not in SIGMA_SPACE_EXCLUDED_IMAGES,
    ),
)


def apply_editorial_rules(text: str, image_name: str = "") -> tuple[str, dict[str, int]]:
    """Resolve editorial tokens; return the text and the per-rule substitution counts.

    A rule with an image-name prefix only applies to rows whose crop name
    starts with it; one with a predicate only where the predicate holds. Only when a rule fires are runs of spaces collapsed and
    the ends stripped, so untouched rows stay byte-identical.
    """
    counts: dict[str, int] = {}
    for name, pattern, replacement, rows in EDITORIAL_RULES:
        if isinstance(rows, str) and not image_name.startswith(rows):
            continue
        if callable(rows) and not rows(image_name):
            continue
        text, count = pattern.subn(replacement, text)
        if count:
            counts[name] = count
    if counts:
        text = MULTIPLE_SPACES.sub(" ", text).strip()
    return text, counts


def normalize_editorial_counts(text: str) -> str:
    """Drop editorial letter counts such as ``[10]``.

    Stavronikita GT marks an illegible stretch with the number of missing
    letters in square brackets. That number is not ink on the page, so it is
    removed; the spaces it leaves are collapsed and the ends stripped.
    """
    if not EDITORIAL_LETTER_COUNT.search(text):
        return text
    return MULTIPLE_SPACES.sub(" ", EDITORIAL_LETTER_COUNT.sub("", text)).strip()


def normalize_greek_gt(text: str) -> str:
    """Apply the agreed Greek GT normalizations."""
    return normalize_dash(
        normalize_colon(
            normalize_question_mark(
                normalize_four_dots(
                    normalize_apostrophe_ticks(
                        normalize_high_commas(
                            normalize_raised_dots(
                                normalize_stigma(normalize_latin_lookalikes(text))
                            )
                        )
                    )
                )
            )
        )
    )


def normalize_text(text: str, image_name: str = "") -> str:
    """Normalize one Greek manifest transcription: NFC, then ``EDITORIAL_RULES``.

    Used by ``resplit.py`` on manifests that already went through
    ``normalize_greek_gt``; the NFC here deliberately supersedes that
    function's non-NFC punctuation choices.
    """
    return apply_editorial_rules(unicodedata.normalize("NFC", text), image_name)[0]


@dataclass(frozen=True)
class ImportRow:
    split: str
    output_name: str
    text: str
    source_image: Path


def normalize_split(value: object) -> str:
    """Normalize source split names to the TrOCR manifest convention."""
    split = str(value).strip().lower()
    aliases = {"validation": "val", "valid": "val"}
    split = aliases.get(split, split)
    if split not in SPLITS:
        raise ValueError(f"Unsupported dataset split: {value!r}")
    return split


def read_esteban_rows(root: Path) -> list[ImportRow]:
    """Read Esteban image names, transcriptions, and splits from its workbook."""
    workbook_path = root / "gray_labels.xlsx"
    image_root = root / "all_bin"
    if not workbook_path.is_file() or not image_root.is_dir():
        raise FileNotFoundError(f"Incomplete Esteban dataset under {root}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    values = worksheet.iter_rows(values_only=True)
    try:
        header = next(values)
    except StopIteration as exc:
        raise ValueError(f"Empty Esteban workbook: {workbook_path}") from exc

    columns = {str(value).strip(): index for index, value in enumerate(header) if value}
    required = {"preproc_file_name", "label", "split"}
    missing = required - columns.keys()
    if missing:
        raise ValueError(f"Missing Esteban workbook columns: {sorted(missing)}")

    rows: list[ImportRow] = []
    seen_names: set[str] = set()
    for row_number, values_row in enumerate(values, start=2):
        image_value = values_row[columns["preproc_file_name"]]
        label_value = values_row[columns["label"]]
        split_value = values_row[columns["split"]]
        if image_value is None and label_value is None and split_value is None:
            continue
        if image_value is None or label_value is None or split_value is None:
            raise ValueError(f"Incomplete Esteban row {row_number}")

        source_name = str(image_value).strip()
        output_name = f"esteban__{source_name}"
        if output_name in seen_names:
            raise ValueError(f"Duplicate Esteban image in workbook: {source_name}")
        seen_names.add(output_name)

        source_image = image_root / source_name
        if not source_image.is_file():
            raise FileNotFoundError(f"Missing Esteban image: {source_image}")
        rows.append(
            ImportRow(
                split=normalize_split(split_value),
                output_name=output_name,
                text=normalize_greek_gt(str(label_value).strip()),
                source_image=source_image,
            )
        )
    workbook.close()
    return rows


def read_text(path: Path) -> str:
    """Read a ground-truth file, including legacy UTF-16 exports."""
    try:
        return path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-16").strip()


def read_labelled_rows(root: Path) -> list[ImportRow]:
    """Read Kraken-style image and ``.gt.txt`` pairs from labelledData."""
    rows: list[ImportRow] = []
    for split in SPLITS:
        image_root = root / "images" / split
        label_root = root / "labels" / split
        if not image_root.is_dir() or not label_root.is_dir():
            raise FileNotFoundError(f"Incomplete labelledData split: {root}/{split}")

        image_paths = sorted(
            path
            for path in image_root.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
        )
        for source_image in image_paths:
            label_path = label_root / f"{source_image.stem}.gt.txt"
            if not label_path.is_file():
                raise FileNotFoundError(f"Missing label for {source_image}: {label_path}")
            rows.append(
                ImportRow(
                    split=split,
                    output_name=f"labelled__{source_image.name}",
                    text=normalize_greek_gt(read_text(label_path)),
                    source_image=source_image,
                )
            )
    return rows


def resplit_import_rows(rows: list[ImportRow], seed: int) -> list[ImportRow]:
    """Pool a source dataset and assign a deterministic 80/10/10 split."""
    shuffled = sorted(rows, key=lambda row: row.output_name)
    random.Random(seed).shuffle(shuffled)
    train_end = round(len(shuffled) * 0.8)
    val_end = round(len(shuffled) * 0.9)
    split_rows: list[ImportRow] = []
    for index, row in enumerate(shuffled):
        split = "train" if index < train_end else "val" if index < val_end else "test"
        split_rows.append(
            ImportRow(
                split=split,
                output_name=row.output_name,
                text=row.text,
                source_image=row.source_image,
            )
        )
    return split_rows


def parse_manifest(path: Path) -> list[tuple[str, str]]:
    """Read one TrOCR filename/transcription manifest."""
    rows: list[tuple[str, str]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line:
            continue
        try:
            image_name, text = line.split("\t", 1)
        except ValueError as exc:
            raise ValueError(f"Invalid manifest row {path}:{line_number}") from exc
        rows.append((image_name, text))
    return rows


def hardlink_or_copy(source: Path, destination: Path) -> None:
    """Hardlink an image when possible, falling back to a copy."""
    try:
        os.link(source, destination)
    except OSError:
        shutil.copy2(source, destination)


def write_manifest(path: Path, rows: list[tuple[str, str]]) -> None:
    """Write one TrOCR manifest."""
    with path.open("w", encoding="utf-8", newline="\n") as output:
        for image_name, text in rows:
            output.write(f"{image_name}\t{text}\n")


def rewrite_greek_manifests() -> dict[str, int]:
    """Rewrite existing Greek GT and copy the new text into combined manifests."""
    from .trocr_splits import rebuild_combined_partition

    summary: dict[str, int] = {}
    for partition in ("pretraining", "finetuning"):
        greek_root = PROCESSED_ROOT / "greek" / partition
        changed = 0
        greek_text: dict[str, str] = {}
        for split in SPLITS:
            path = greek_root / f"gt_{split}.txt"
            if not path.is_file():
                continue
            rows = []
            for image_name, text in parse_manifest(path):
                cleaned = normalize_greek_gt(text)
                if cleaned != text:
                    changed += 1
                greek_text[image_name] = cleaned
                rows.append((image_name, cleaned))
            write_manifest(path, rows)
        summary[f"greek_{partition}_rows_changed"] = changed
        try:
            rebuild_combined_partition(partition)
            summary[f"combined_{partition}"] = 1
        except FileNotFoundError:
            combined_root = PROCESSED_ROOT / "combined" / partition
            synced = 0
            for split in SPLITS:
                path = combined_root / f"gt_{split}.txt"
                if not path.is_file():
                    continue
                rows = []
                for image_name, text in parse_manifest(path):
                    updated = greek_text.get(image_name, text)
                    if updated != text:
                        synced += 1
                    rows.append((image_name, updated))
                write_manifest(path, rows)
            summary[f"combined_{partition}_rows_synced"] = synced
    return summary


def stage_partition(target: Path, imports: list[ImportRow]) -> tuple[Path, dict[str, int]]:
    """Build a validated replacement for one pretraining partition."""
    if not target.is_dir():
        raise FileNotFoundError(f"Missing TrOCR pretraining partition: {target}")

    staging = target.with_name(f".{target.parent.name}_pretraining_staging")
    if staging.exists():
        shutil.rmtree(staging)
    image_dir = staging / "image"
    image_dir.mkdir(parents=True)

    imported_by_split = {split: [row for row in imports if row.split == split] for split in SPLITS}
    summary: dict[str, int] = {}
    for split in SPLITS:
        manifest = target / f"gt_{split}.txt"
        if not manifest.is_file():
            raise FileNotFoundError(f"Missing TrOCR manifest: {manifest}")

        rows = [
            (image_name, text)
            for image_name, text in parse_manifest(manifest)
            if not image_name.startswith(IMPORTED_PREFIXES)
        ]
        seen_names = {image_name for image_name, _ in rows}
        if len(seen_names) != len(rows):
            raise ValueError(f"Duplicate existing rows in {manifest}")

        for image_name, _ in rows:
            source_image = target / "image" / image_name
            if not source_image.is_file():
                raise FileNotFoundError(f"Missing existing TrOCR image: {source_image}")
            hardlink_or_copy(source_image, image_dir / image_name)

        for import_row in imported_by_split[split]:
            if import_row.output_name in seen_names:
                raise ValueError(f"Duplicate imported image name: {import_row.output_name}")
            seen_names.add(import_row.output_name)
            hardlink_or_copy(
                import_row.source_image,
                image_dir / import_row.output_name,
            )
            rows.append((import_row.output_name, import_row.text))

        write_manifest(staging / f"gt_{split}.txt", rows)
        summary[split] = len(rows)
    return staging, summary


def replace_partitions(staged_targets: list[tuple[Path, Path]]) -> None:
    """Swap staged partitions into place, rolling both back on failure."""
    backups: list[tuple[Path, Path]] = []
    try:
        for target, staging in staged_targets:
            backup = target.with_name(f".{target.parent.name}_pretraining_backup")
            if backup.exists():
                shutil.rmtree(backup)
            target.rename(backup)
            backups.append((target, backup))
            staging.rename(target)
    except Exception:
        for target, backup in reversed(backups):
            if target.exists():
                shutil.rmtree(target)
            backup.rename(target)
        raise
    else:
        for _, backup in backups:
            shutil.rmtree(backup)


def main() -> None:
    """Import both Greek datasets into Greek and combined pretraining only."""
    esteban_rows = resplit_import_rows(
        read_esteban_rows(ESTEBAN_ROOT),
        SPLIT_SEED,
    )
    labelled_rows = resplit_import_rows(
        read_labelled_rows(LABELLED_ROOT),
        SPLIT_SEED + 1,
    )
    imports = esteban_rows + labelled_rows

    staged_targets: list[tuple[Path, Path]] = []
    target_summaries: dict[str, dict[str, int]] = {}
    try:
        for target in TARGETS:
            staging, summary = stage_partition(target, imports)
            staged_targets.append((target, staging))
            target_summaries[str(target.relative_to(PROCESSED_ROOT))] = summary
        replace_partitions(staged_targets)
    except Exception:
        for _, staging in staged_targets:
            if staging.exists():
                shutil.rmtree(staging)
        raise

    source_summary = {
        "esteban": {split: sum(row.split == split for row in esteban_rows) for split in SPLITS},
        "labelled": {split: sum(row.split == split for row in labelled_rows) for split in SPLITS},
    }
    print(
        json.dumps(
            {
                "sources": source_summary,
                "updated_pretraining_partitions": target_summaries,
                "finetuning_partitions_changed": False,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
