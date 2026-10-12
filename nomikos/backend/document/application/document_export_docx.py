"""The document's transcription as a Word file, one heading per page.

The plain text export gives the words and nothing else; this gives the same words in
the shape an editor works in. Word and LibreOffice need two things a ``.txt`` cannot
say: which font to draw a line in, and which way it runs. Without them a Syriac line
opens left to right in a Latin font, which is either boxes or a mirror image of the
manuscript. Both come from the line's own characters (see ``script_detection``), so a
Greek gloss on a Syriac page stays left to right.

Same pages, same text and same page numbers as the text export: the selector, the
ground-truth rule and the position-based numbering are shared, not restated. So is the
region grouping (ADR 0011): a labelled region's lines follow its label in bold.
"""

from __future__ import annotations

import asyncio
import io
from dataclasses import dataclass
from uuid import UUID

from docx import Document as new_word_document
from docx.document import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from sqlalchemy.ext.asyncio import AsyncSession

from backend.annotation.application.page_xml_export_service import export_file_stem
from backend.core.exceptions import NotFoundError
from backend.core.fonts import SCRIPT_FONT_FAMILIES
from backend.document.application.document_export_selection import (
    DocumentExportSelector,
    document_download_name,
)
from backend.document.application.document_export_text import _ground_truth_text
from backend.document.application.document_service import DocumentService
from backend.document.application.regions import RegionHeading, region_text
from backend.document.application.script_detection import Script, detect_script, is_rtl
from backend.document.infrastructure.document_repository import DocumentRepository
from backend.users.infrastructure.orm_models import User

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_BODY_FONT_SIZE = Pt(12)
# C0 controls other than tab, newline and carriage return (which python-docx turns into
# w:tab and w:br), and the two Unicode non-characters XML 1.0 forbids.
_XML_INVALID = dict.fromkeys(
    [*(c for c in range(0x20) if c not in (0x09, 0x0A, 0x0D)), 0xFFFE, 0xFFFF]
)


@dataclass(frozen=True)
class DocumentDocxExport:
    filename: str
    content: bytes


@dataclass(frozen=True)
class DocxPage:
    """One page reduced to its number and its transcribed lines, off the session.

    A :class:`RegionHeading` among the lines opens a labelled region.
    """

    number: int
    lines: list[str | RegionHeading]


class DocumentDocxExportService:
    def __init__(
        self,
        *,
        documents: DocumentRepository | None = None,
        selector: DocumentExportSelector | None = None,
        document_service: DocumentService | None = None,
    ) -> None:
        self._documents = documents or DocumentRepository()
        self._selector = selector or DocumentExportSelector(documents=self._documents)
        self._document_service = document_service or DocumentService()

    async def export_document_docx(
        self,
        session: AsyncSession,
        user: User,
        project_id: UUID,
        document_id: UUID,
        *,
        reviewed_only: bool = False,
    ) -> DocumentDocxExport:
        selection = await self._selector.select(
            session, user, project_id, document_id, reviewed_only=reviewed_only
        )
        pages: list[DocxPage] = []
        for part in selection.parts:
            lines = await self._documents.list_part_lines(session, part.id)
            blocks = await self._documents.list_part_blocks(session, part.id)
            pages.append(
                DocxPage(number=selection.page_number(part), lines=_page_text(blocks, lines))
            )
        content = await asyncio.to_thread(render_docx, pages)
        return DocumentDocxExport(
            filename=document_download_name(selection.document.name, "transcription.docx"),
            content=content,
        )

    async def export_part_docx(
        self,
        session: AsyncSession,
        user: User,
        project_id: UUID,
        document_id: UUID,
        part_id: UUID,
    ) -> DocumentDocxExport:
        """One page as a Word file, for a project member in the page editor."""
        document = await self._document_service.get_document(session, user, project_id, document_id)
        part = await self._documents.get_part(session, part_id)
        if part is None or part.document_id != document.id:
            raise NotFoundError("Part not found")
        page = await self._part_page(session, part)
        content = await asyncio.to_thread(render_docx, [page])
        return DocumentDocxExport(
            filename=f"{export_file_stem(document.name, page.number)}.docx",
            content=content,
        )

    async def export_part_docx_public(
        self,
        session: AsyncSession,
        project_id: UUID,
        document_id: UUID,
        part_id: UUID,
        *,
        token: str | None = None,
    ) -> DocumentDocxExport:
        """One published page as a Word file, for the anonymous reader.

        The public surface exports page by page, as its PDF and PAGE XML do, so this is
        held to the published-part gate those use rather than to document membership.
        """
        context = await self._document_service.get_published_part_context(
            session, project_id, document_id, part_id, token=token
        )
        page = await self._part_page(session, context.part)
        content = await asyncio.to_thread(render_docx, [page])
        return DocumentDocxExport(
            filename=f"{export_file_stem(context.document.name, page.number)}.docx",
            content=content,
        )

    async def _part_page(self, session: AsyncSession, part) -> DocxPage:
        number = await self._documents.part_page_number(session, part)
        lines = await self._documents.list_part_lines(session, part.id)
        blocks = await self._documents.list_part_blocks(session, part.id)
        return DocxPage(number=number, lines=_page_text(blocks, lines))


def _page_text(blocks, lines) -> list[str | RegionHeading]:
    return region_text(blocks, lines, _ground_truth_text)


def render_docx(pages: list[DocxPage]) -> bytes:
    """Write ``pages`` as a Word document: a heading per page, a paragraph per line."""
    word = new_word_document()
    _set_default_font(word)
    for index, page in enumerate(pages):
        if index > 0:
            # A reader continuing onto the next page should find it at the top of a
            # fresh sheet, not where the previous page's last line happened to end.
            word.add_page_break()
        # A page with nothing transcribed still gets its heading, for the reason the
        # text export keeps its marker: a missing "Page 7" reads as a missing page.
        word.add_heading(f"Page {page.number}", level=2)
        for item in page.lines:
            if isinstance(item, RegionHeading):
                _add_region_heading(word, item.label)
            else:
                _add_line(word, item)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _set_default_font(word: WordDocument) -> None:
    normal = word.styles["Normal"]
    normal.font.size = _BODY_FONT_SIZE
    run_properties = normal.element.get_or_add_rPr()
    _set_fonts(run_properties, SCRIPT_FONT_FAMILIES[Script.latin.value])
    # Word sizes complex-script runs (Syriac among them) by ``szCs``, not ``sz``.
    _add_in_schema_order(run_properties, "w:szCs").set(
        qn("w:val"), str(int(_BODY_FONT_SIZE.pt * 2))
    )


def _xml_safe(text: str) -> str:
    """Drop the characters XML cannot hold, which python-docx refuses with a ValueError.

    Approved text is not sanitised on the way in, and Word puts a vertical tab on the
    clipboard for a manual line break, so one pasted line would fail the whole export.
    The vertical tab becomes a newline, which python-docx writes as a line break.
    """
    return text.replace("\x0b", "\n").translate(_XML_INVALID)


def _add_region_heading(word: WordDocument, label: str) -> Paragraph:
    """A region's label on a line of its own, in bold: a folio name, not a heading level."""
    paragraph = word.add_paragraph()
    paragraph.add_run(_xml_safe(label)).bold = True
    return paragraph


def _add_line(word: WordDocument, text: str) -> Paragraph:
    text = _xml_safe(text)
    script = detect_script(text)
    paragraph = word.add_paragraph()
    run = paragraph.add_run(text)
    run_properties = run._r.get_or_add_rPr()
    _set_fonts(run_properties, SCRIPT_FONT_FAMILIES[script.value])
    if is_rtl(script):
        # ``w:bidi`` makes the paragraph right to left, which also makes its natural
        # alignment the right margin; ``w:rtl`` tells Word the run's characters are a
        # right-to-left script to be laid out and shaped as one.
        _add_in_schema_order(paragraph._p.get_or_add_pPr(), "w:bidi")
        run.font.rtl = True
    return paragraph


def _set_fonts(run_properties, family: str) -> None:
    """Name ``family`` for every slot Word picks a font from.

    ``ascii`` and ``hAnsi`` cover Latin and Greek, ``cs`` covers complex scripts such as
    Syriac, and ``eastAsia`` is set too so no slot is left to a theme font that lacks
    the glyphs.
    """
    fonts = run_properties.get_or_add_rFonts()
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(slot), family)


# What may follow each element we add, in OOXML schema order (ECMA-376 part 1, CT_RPr
# and CT_PPrBase). Word refuses a file whose property children are out of order, and
# python-docx has no setter for ``szCs`` or ``bidi``.
_SUCCESSORS = {
    "w:szCs": (
        "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
        "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
    ),
    "w:bidi": (
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
        "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange",
    ),
}  # fmt: skip


def _add_in_schema_order(parent, tag: str):
    """Add ``tag`` to ``parent`` before any sibling the schema places after it."""
    return parent.insert_element_before(OxmlElement(tag), *_SUCCESSORS[tag])
